"""Build the presentation (docs/Project_Presentation.pptx), 15 slides, native charts.

Numbers come from benchmark/results/*.json and imgframe/profiles.json.

    python docs/build_slides.py
"""
import json
import statistics as st
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "docs" / "figures"
RES = ROOT / "benchmark" / "results"

NAVY = RGBColor(0x0F, 0x2A, 0x44)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x64, 0x74, 0x8B)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_L = RGBColor(0xCC, 0xFB, 0xF1)
AMBER = RGBColor(0xD9, 0x77, 0x06)
GREY = RGBColor(0x94, 0xA3, 0xB8)
PAPER = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
HEAD, BODY = "Cambria", "Calibri"
W, H = Inches(13.333), Inches(7.5)

TEAM = [("Prapti Kurkuri", "24BVD0185"), ("Vaanya Garg", "24BVD0139"),
        ("Aditya Manoj", "24BVD0070"), ("Arrnav Trisal", "24BVD0064")]


def load(name):
    return json.loads((RES / name).read_text())


# ------------------------------------------------------------------ helpers
def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None, italic=False):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, _, rest = item.partition("|")
        r0 = p.add_run()
        r0.text = "▸  "
        r0.font.color.rgb = TEAL
        r0.font.size = Pt(size)
        r0.font.name = BODY
        r = p.add_run()
        r.text = head if rest else item
        r.font.bold = bool(rest)
        r.font.size = Pt(size)
        r.font.name = BODY
        r.font.color.rgb = color
        if rest:
            r2 = p.add_run()
            r2.text = " " + rest
            r2.font.size = Pt(size)
            r2.font.name = BODY
            r2.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def badge(slide, x, y, label, fill=TEAL, size=0.62):
    c = rect(slide, x, y, Inches(size), Inches(size), fill, shape=MSO_SHAPE.OVAL, name=f"badge {label}")
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(20 if len(label) < 2 else 15)
    r.font.bold = True
    r.font.name = HEAD
    r.font.color.rgb = WHITE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def slide_base(prs, title, kicker=None, num=None, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else PAPER
    x = Inches(0.6)
    if num:
        badge(s, x, Inches(0.45), num)
        x = Inches(1.4)
    if kicker:
        text(s, x, Inches(0.38), Inches(10), Inches(0.3), kicker.upper(), size=11, color=TEAL, bold=True)
    text(s, x, Inches(0.62), Inches(11.5), Inches(0.7), title, size=28, font=HEAD,
         color=WHITE if dark else NAVY, bold=True, name="Title")
    idx = len(prs.slides)
    if not dark:
        text(s, Inches(12.2), Inches(7.05), Inches(0.8), Inches(0.3), str(idx), size=10, color=MUTED,
             align=PP_ALIGN.RIGHT)
        text(s, Inches(0.6), Inches(7.05), Inches(8), Inches(0.3),
             "Cloud-Native Parallel Image Processing · BECE355L", size=10, color=MUTED)
    return s


def stat(slide, x, y, w, value, label, color=TEAL, vsize=40, lsize=13):
    text(slide, x, y, w, Inches(0.8), value, size=vsize, font=HEAD, color=color, bold=True)
    text(slide, x, y + Inches(0.85), w, Inches(0.8), label, size=lsize, color=MUTED)


def card(slide, x, y, w, h, title, body, accent=TEAL, size=14):
    rect(slide, x, y, w, h, WHITE, line=LINE, name=f"card {title}")
    rect(slide, x, y, Inches(0.07), h, accent)
    text(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.4), Inches(0.4), title, size=16,
         font=HEAD, bold=True, color=NAVY)
    text(slide, x + Inches(0.25), y + Inches(0.62), w - Inches(0.4), h - Inches(0.7), body, size=size,
         color=INK)


def picture(slide, path, x, y, max_w, max_h):
    from PIL import Image
    with Image.open(path) as im:
        pw, ph = im.size
    scale = min(max_w / pw, max_h / ph)
    w, h = int(pw * scale), int(ph * scale)
    return slide.shapes.add_picture(str(path), x + (max_w - w) // 2, y + (max_h - h) // 2, w, h)


def style_chart(chart, title=None, legend=True, size=12, number_format=None, labels=True):
    chart.font.size = Pt(size)
    chart.font.name = BODY
    chart.font.color.rgb = INK
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = title
        tp = chart.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size = Pt(size + 1)
        tp.runs[0].font.bold = True
        tp.runs[0].font.color.rgb = NAVY
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
    try:
        va = chart.value_axis
        va.major_gridlines.format.line.color.rgb = LINE
        va.format.line.fill.background()
        va.tick_labels.font.color.rgb = MUTED
        va.tick_labels.font.size = Pt(size - 1)
        if number_format:
            va.tick_labels.number_format = number_format
            va.tick_labels.number_format_is_linked = False
        ca = chart.category_axis
        ca.format.line.color.rgb = GREY
        ca.tick_labels.font.size = Pt(size - 1)
        ca.tick_labels.font.color.rgb = INK
        ca.has_major_gridlines = False
    except Exception:
        pass
    if labels:
        pl = chart.plots[0]
        pl.has_data_labels = True
        dl = pl.data_labels
        dl.font.size = Pt(size - 2)
        dl.font.color.rgb = INK
        if number_format:
            dl.number_format = number_format
            dl.number_format_is_linked = False


def color_series(chart, colors):
    for s, c in zip(chart.plots[0].series, colors):
        f = s.format
        if chart.chart_type in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
            f.line.color.rgb = c
            f.line.width = Pt(2.5)
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = c
            s.marker.format.line.color.rgb = c
            s.smooth = False
        else:
            f.fill.solid()
            f.fill.fore_color.rgb = c


def autoscale_summary():
    d = RES / "autoscale"
    runs = {}
    for f in sorted(d.glob("*_run*.json")) if d.exists() else []:
        s = json.loads(f.read_text())["summary"]
        runs.setdefault(s["label"], []).append(s)
    return {k: {m: st.mean([r[m] for r in v if r[m] is not None] or [0])
                for m in ("p50_s", "p95_s", "within_target_pct", "pod_seconds", "first_scale_up_s",
                          "scaling_events")} | {"runs": len(v), "jobs": v[0]["jobs"], "target": v[0]["target_s"]}
            for k, v in runs.items()}


# ------------------------------------------------------------------- slides
def build() -> Path:
    eng = load("planner_engine.json")
    clu = load("planner_cluster.json")
    sto = load("storage_study.json")
    halo = load("halo_calibration.json")
    prof = json.loads((ROOT / "imgframe" / "profiles.json").read_text())
    auto = autoscale_summary()
    rm, sm = sto[0], sto[1]
    er = {r["op"]: r for r in eng["rows"]}

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 ── title
    s = slide_base(prs, "", dark=True)
    text(s, Inches(0.8), Inches(1.0), Inches(11), Inches(0.4), "BECE355L · ADVANCED CLOUD COMPUTING · FALL 2026–27",
         size=13, color=TEAL, bold=True)
    text(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(2.0),
         "Cloud-Native Parallel Image Processing", size=46, font=HEAD, color=WHITE, bold=True)
    text(s, Inches(0.8), Inches(3.35), Inches(11.5), Inches(1.0),
         "Object storage · deep-learning operations · a cost-aware planner · work-aware autoscaling",
         size=22, color=RGBColor(0xCB, 0xD5, 0xE1), italic=True)
    for i, (n, reg) in enumerate(TEAM):
        x = Inches(0.8 + 3.0 * i)
        rect(s, x, Inches(4.7), Inches(0.06), Inches(0.75), TEAL)
        text(s, x + Inches(0.2), Inches(4.7), Inches(2.7), Inches(0.4), n, size=17, color=WHITE, bold=True)
        text(s, x + Inches(0.2), Inches(5.1), Inches(2.7), Inches(0.4), reg, size=13,
             color=RGBColor(0xCB, 0xD5, 0xE1))
    text(s, Inches(0.8), Inches(6.3), Inches(11), Inches(0.4),
         "Faculty: Aparna Mohanty  ·  SENSE, VIT Vellore", size=14, color=RGBColor(0xCB, 0xD5, 0xE1))
    s.notes_slide.notes_text_frame.text = (
        "Introduce the team. One line: we took our tile-parallel image framework and made it cost-aware: "
        "where data lives, which operations it can run, how it plans a job and how it scales.")

    # 2 ── where v1 stood
    s = slide_base(prs, "Version 1 worked — and showed four limits", "Starting point")
    text(s, Inches(0.6), Inches(1.45), Inches(12), Inches(0.8),
         "Image → halo tiles → one Redis task per tile → stateless worker pods (leases, retries) → stitched "
         "output identical to whole-image processing.", size=16, color=MUTED)
    lim = [(f"{rm['redis_bytes_per_8192_job'] / 1e6:.0f} MB", "of Redis per 8192² job:\nevery tile byte went through Redis", AMBER),
           ("2 jobs", "fit in 1 GB of Redis;\n12288² images were rejected", AMBER),
           ("0.45×", "Canny when distributed:\nmoving a tile cost more than filtering it", AMBER),
           ("1 tile = 1", "autoscaler counted tiles,\nbut an HED tile costs ~1800× a Canny tile", AMBER)]
    for i, (v, lab, c) in enumerate(lim):
        x = Inches(0.6 + 3.1 * i)
        rect(s, x, Inches(2.55), Inches(2.85), Inches(2.6), WHITE, line=LINE)
        stat(s, x + Inches(0.25), Inches(2.8), Inches(2.5), v, lab, color=c, vsize=36, lsize=13)
    text(s, Inches(0.6), Inches(5.6), Inches(12), Inches(1.2), [
        [("Four upgrades, one per limit:  ", {"bold": True, "color": NAVY}),
         ("①  object storage   ②  deep-learning ops   ③  cost-aware planner   ④  work-aware autoscaling",
          {"color": INK})]], size=17)
    s.notes_slide.notes_text_frame.text = (
        "v1: FastAPI partitions, Redis reliable queue with Lua scripts, leases and heartbeats, last finisher "
        "stitches. Bilateral reached 4x, but four things limited it: memory, cheap ops, manual config, "
        "tile-count scaling.")

    # 3 ── architecture
    s = slide_base(prs, "Architecture after the upgrades", "System")
    picture(s, FIG / "fig_architecture_v2.png", Inches(0.5), Inches(1.45), Inches(8.3), Inches(5.45))
    bullets(s, Inches(9.1), Inches(1.6), Inches(3.8), Inches(5.3), [
        "Bytes in S3 / MinIO:|input, tiles, results, output",
        "Redis = control plane:|task IDs, leases, costs, work-seconds",
        "API plans every job:|local or distributed, tile size, workers",
        "Workers:|OpenCV or ONNX Runtime, stateless",
        "One image,|three ways to run: local · Compose · Kubernetes",
    ], size=15, gap=10)
    s.notes_slide.notes_text_frame.text = (
        "Walk the data flow: upload, plan, tiles to bucket, IDs to Redis, workers claim with lease, GET tile, "
        "process, PUT result, complete via Lua, last finisher stitches, client gets a presigned URL.")

    # 4 ── upgrade 1 design
    s = slide_base(prs, "Bytes to object storage, IDs to Redis", "Upgrade 1 · Object storage", "1")
    card(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(2.45), "Claim-check data plane",
         "Tiles go to jobs/<id>/tiles/ before any task is queued; a task is just an ID.\n"
         "Same boto3 code for MinIO (local) and Amazon S3.\n"
         "Results: 307 redirect to a presigned URL, bytes never pass the API.\n"
         "Outputs > 150 MP stream into a tiled BigTIFF.")
    card(s, Inches(0.6), Inches(4.25), Inches(6.0), Inches(2.45), "Security by design (AWS templates)",
         "Least-privilege IAM: API puts tiles; workers get tiles, put results.\n"
         "SSE-KMS, Block Public Access, TLS-only bucket policy.\n"
         "IRSA: pods get roles, no keys in containers.\n"
         "Lifecycle rule expires job data after 1 day.", accent=NAVY)
    stat(s, Inches(7.3), Inches(1.7), Inches(5.5), f"{rm['redis_bytes_per_8192_job'] / 1e6:.0f} MB → "
         f"{sm['redis_bytes_per_8192_job'] / 1e3:.0f} KB", "Redis memory per 8192² job", vsize=36)
    stat(s, Inches(7.3), Inches(3.4), Inches(5.5), f"{rm['concurrent_8192_jobs_accepted']} → "
         f"{sm['concurrent_8192_jobs_accepted']}+", "concurrent 8192² jobs in 1 GB Redis (all 12 tried fit)",
         vsize=36)
    stat(s, Inches(7.3), Inches(5.1), Inches(5.5), f"+{sm['io_ms_per_tile'] - rm['io_ms_per_tile']:.1f} ms",
         f"I/O per tile — the price ({rm['io_ms_per_tile']:.1f} → {sm['io_ms_per_tile']:.1f} ms)", color=AMBER,
         vsize=36)
    s.notes_slide.notes_text_frame.text = (
        "Redis footprint per job shrinks ~4000x. The trade-off: each tile now costs an S3 PUT and GET; a small "
        "bilateral job went from 0.82 to 1.09 s. The planner knows this overhead. AWS templates were written "
        "but not deployed - MinIO stood in for S3.")

    # 5 ── upgrade 2: neural ops
    s = slide_base(prs, "Deep-learning operations: HED edges and Real-ESRGAN ×4", "Upgrade 2 · Neural ops", "2")
    picture(s, FIG / "fig_canny_vs_hed.png", Inches(0.5), Inches(1.5), Inches(12.3), Inches(3.3))
    bullets(s, Inches(0.6), Inches(5.05), Inches(12.2), Inches(1.9), [
        "HED (VGG-16, 5 side outputs):|we wrote a Caffe→ONNX converter because OpenCV 5 dropped Caffe import",
        "Real-ESRGAN realesr-general-x4v3:|×4 output, so the stitcher scales tile coordinates",
        "ONNX Runtime, 1 thread per process:|all parallelism comes from the framework, a fair comparison",
    ], size=15, gap=6)
    s.notes_slide.notes_text_frame.text = (
        "Canny gives thin noisy edges, HED gives semantic object contours. The error maps on the right show "
        "tiled minus whole-image output with the calibrated halo: essentially zero.")

    # 6 ── halo calibration
    s = slide_base(prs, "Calibrating the halo from the effective receptive field", "Upgrade 2 · Halo", "2")
    cd = CategoryChartData()
    hs = [x for x in halo["hed"]["sweep"]]
    cd.categories = [str(x["halo"]) for x in hs]
    cd.add_series("HED seam PSNR (dB; bit-exact shown as 70)", [min(x["seam_psnr_db"], 70) for x in hs])
    cd.add_series("40 dB threshold", [40] * len(hs))
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(1.5), Inches(7.2), Inches(5.3), cd)
    ch = gf.chart
    style_chart(ch, "HED: seam PSNR vs halo (px), 256 px tiles", labels=False)
    color_series(ch, [TEAL, AMBER])
    ch.category_axis.has_title = True
    ch.category_axis.axis_title.text_frame.text = "halo (px)"
    ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    hp, hf = (next(x for x in hs if x["halo"] == v) for v in (halo["hed"]["halo"], halo["hed"]["theoretical_halo"]))
    ss = halo["superres"]["sweep"]
    sp, sf = (next(x for x in ss if x["halo"] == v)
              for v in (halo["superres"]["halo"], halo["superres"]["theoretical_halo"]))
    stat(s, Inches(8.1), Inches(1.6), Inches(4.8), f"98 → {halo['hed']['halo']} px",
         f"HED halo: extra compute {100 * hf['overhead_theory']:.0f} % → {100 * hp['overhead_theory']:.0f} %",
         vsize=34)
    stat(s, Inches(8.1), Inches(3.15), Inches(4.8), f"34 → {halo['superres']['halo']} px",
         f"Real-ESRGAN halo: extra compute {100 * sf['overhead_theory']:.0f} % → "
         f"{100 * sp['overhead_theory']:.0f} %", vsize=34)
    bullets(s, Inches(8.1), Inches(4.75), Inches(4.8), Inches(2.2), [
        "Criterion:|seam PSNR ≥ 40 dB and p99.9 error ≤ 8 grey levels",
        "Stride fix:|snap halo edges to 16 px, theoretical halo became bit-exact",
    ], size=14, gap=6)
    s.notes_slide.notes_text_frame.text = (
        "Theoretical receptive field is exact but expensive. Luo et al. showed the effective receptive field "
        "is much smaller. We sweep the halo and take the smallest that meets the criterion. The 16-px snap "
        "fixed a pooling-grid misalignment.")

    # 7 ── neural scaling reality
    s = slide_base(prs, "Neural ops are memory-bound on one host", "Measured capacity curve S(n)")
    cap = prof["capacity"]
    ns = ["1", "2", "4", "6", "8", "12", "16"]
    cd = CategoryChartData()
    cd.categories = ns
    for op, lab in (("bilateral", "Bilateral"), ("hed", "HED"), ("superres", "Real-ESRGAN")):
        cd.add_series(lab, [cap[op][n] for n in ns])
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(1.5), Inches(7.6), Inches(5.35), cd)
    ch = gf.chart
    style_chart(ch, "Throughput of n busy workers relative to one", labels=False, number_format='0.0"×"')
    color_series(ch, [NAVY, TEAL, AMBER])
    ch.category_axis.has_title = True
    ch.category_axis.axis_title.text_frame.text = "concurrent workers"
    ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    stat(s, Inches(8.6), Inches(1.6), Inches(4.3), f"{cap['hed']['8']:.1f}× at 8", "HED peak, then it falls")
    stat(s, Inches(8.6), Inches(3.05), Inches(4.3), f"≈ {cap['superres']['4']:.1f}× from 4",
         "Real-ESRGAN plateau", color=AMBER)
    text(s, Inches(8.6), Inches(4.6), Inches(4.3), Inches(2.3),
         "The plan expected near-linear speedup. Measured: activations of many concurrent tiles overflow the "
         "caches, so extra workers wait on DRAM. On a multi-node cluster each node adds bandwidth; the planner "
         "uses the measured curve either way.", size=14, color=INK)
    s.notes_slide.notes_text_frame.text = (
        "Important honest result: the neural ops do not scale linearly on a single 8-core laptop. Bilateral is "
        "compute-bound and keeps scaling to ~7x. This is why the planner needs a capacity curve.")

    # 8 ── planner model
    s = slide_base(prs, "A calibrated cost model decides how to run each job", "Upgrade 3 · Planner", "3")
    rect(s, Inches(0.6), Inches(1.6), Inches(7.2), Inches(2.3), WHITE, line=LINE)
    text(s, Inches(0.85), Inches(1.75), Inches(6.8), Inches(2.1), [
        [("T", {"italic": True}), ("dist", {"size": 13}), (" = T", {"italic": True}), ("split", {"size": 13}),
         (" + ⌈K/n⌉ · c · p̄ · n / S(n) + overhead + T", {"italic": True}), ("stitch", {"size": 13})],
        [("o(p) = o", {"italic": True}), ("0", {"size": 13}), (" + o", {"italic": True}), ("px", {"size": 13}),
         (" · p", {"italic": True}), ("halo", {"size": 13}), (" · (1 + s²)", {"italic": True})],
        [("T", {"italic": True}), ("local", {"size": 13}), (" = c · P", {"italic": True})],
    ], size=20, font=HEAD, color=NAVY)
    bullets(s, Inches(0.6), Inches(4.2), Inches(7.2), Inches(2.7), [
        "c:|measured cost per pixel of each op (0.3 ns grayscale … 24 µs Real-ESRGAN)",
        "S(n):|measured capacity curve per op",
        "Search:|tile ∈ {256, 512, 1024, 2048} × every worker count; 20 % margin; near-ties → fewer workers",
        "Online:|workers update c with a contention-normalised EMA in Redis",
    ], size=14, gap=6)
    picture(s, FIG / "fig_capacity.png", Inches(8.1), Inches(1.6), Inches(4.8), Inches(5.2))
    s.notes_slide.notes_text_frame.text = (
        "Ernest-style: calibrate a parametric model with a few measurements, then extrapolate. The planner "
        "answers three questions: distribute or not, tile size, how many workers.")

    # 9 ── planner results
    s = slide_base(prs, "The planner removes the slowdowns of naive tiling", "Upgrade 3 · Results", "3")
    ops = ["grayscale", "canny", "gaussian_blur", "sobel", "bilateral", "hed", "superres"]
    lab = {"grayscale": "Grayscale", "canny": "Canny", "gaussian_blur": "Gaussian", "sobel": "Sobel",
           "bilateral": "Bilateral", "hed": "HED", "superres": "Real-ESRGAN"}
    cd = CategoryChartData()
    cd.categories = [lab[o] for o in ops]
    cd.add_series("Naive: 512 px tiles, 16 workers", [er[o]["naive_speedup"] for o in ops])
    cd.add_series("Planner", [er[o]["plan_speedup"] for o in ops])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(8.4), Inches(5.35), cd)
    ch = gf.chart
    style_chart(ch, "Speedup over sequential (in-memory engine)", number_format='0.00"×"')
    color_series(ch, [GREY, TEAL])
    ch.plots[0].gap_width = 60
    ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    stat(s, Inches(9.3), Inches(1.6), Inches(3.7), f"{er['grayscale']['naive_speedup']:.2f}× → "
         f"{er['grayscale']['plan_speedup']:.2f}×", "grayscale: planner runs it locally", vsize=32)
    stat(s, Inches(9.3), Inches(3.15), Inches(3.7), f"{er['hed']['naive_speedup']:.1f}× → "
         f"{er['hed']['plan_speedup']:.1f}×", "HED: 8 workers beat 16", vsize=32)
    stat(s, Inches(9.3), Inches(4.7), Inches(3.7), f"{eng['mape_pct']:.0f} % vs {eng['mape_simple_pct']:.0f} %",
         "prediction error (MAPE), refined vs linear model", color=NAVY, vsize=32)
    s.notes_slide.notes_text_frame.text = (
        f"Naive tiling made every cheap op slower. Planner keeps them local. Bilateral: chooses 12 workers, "
        f"within 3% of naive. Neural: picks fewer workers and wins. Full framework with S3 (8 workers): "
        f"16 x 1024-px tiles instead of 64 x 512, 4-14% faster for classical ops; cluster MAPE "
        f"{clu['mape_pct']:.0f}%.")

    # 10 ── model accuracy
    s = slide_base(prs, "Modelling contention is what makes predictions useful", "Upgrade 3 · Accuracy", "3")
    mo = eng["mape_by_op"]
    ops2 = ["bilateral", "hed", "superres", "gaussian_blur", "sobel", "canny", "grayscale"]
    cd = CategoryChartData()
    cd.categories = [lab[o] for o in ops2]
    cd.add_series("Linear model", [mo[o]["simple"] / 100 for o in ops2])
    cd.add_series("Refined model (capacity curve + serial dispatch)", [mo[o]["refined"] / 100 for o in ops2])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.5), Inches(7.8), Inches(5.35), cd)
    ch = gf.chart
    style_chart(ch, "Mean absolute % error per operation (lower is better)", number_format="0%")
    color_series(ch, [GREY, TEAL])
    ch.plots[0].gap_width = 50
    ch.category_axis.reverse_order = True
    picture(s, FIG / "fig_model_accuracy.png", Inches(8.6), Inches(1.5), Inches(4.3), Inches(3.6))
    text(s, Inches(8.6), Inches(5.3), Inches(4.3), Inches(1.6),
         "Biggest gains on the expensive ops, where a wrong plan costs seconds. Cheap jobs (< 0.2 s) stay "
         "noisy, but there a wrong choice costs milliseconds.", size=14)
    s.notes_slide.notes_text_frame.text = (
        "A linear model assumes n workers give n-fold throughput. That is wrong for memory-bound ops, so it "
        "overpredicts speedup; the capacity curve fixes it.")

    # 11 ── work-aware autoscaler
    s = slide_base(prs, "Scale on queued work, not on queue length", "Upgrade 4 · Autoscaling", "4")
    rect(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(1.25), WHITE, line=LINE)
    text(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(1.25), [
        [("n = clamp( max( ⌈W / (D·u)⌉ , ⌈λ / u⌉ ), 1, 8 )", {})]],
        size=26, font=HEAD, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=True)
    card(s, Inches(0.6), Inches(3.15), Inches(3.9), Inches(3.6), "W — queued work",
         "Each tile's predicted compute-seconds is added when queued and removed exactly (in Lua) on completion, "
         "give-up or delete.")
    card(s, Inches(4.7), Inches(3.15), Inches(3.9), Inches(3.6), "D, u, λ",
         "D = latency target 30 s − pod start-up 5 s.\nu = 0.8 target utilisation.\nλ = EMA of the rate "
         "at which work arrives.\nIdle (W ≈ 0) for 10 s → drop straight to 1 worker.", accent=NAVY)
    card(s, Inches(8.8), Inches(3.15), Inches(3.9), Inches(3.6), "Baselines",
         "HPA: CPU 70 % of request, 30 s scale-down window.\nKEDA 2.21: Redis list length, 16 tiles per "
         "replica, 3 s polling.\nSame 1–8 replicas for all.", accent=AMBER)
    s.notes_slide.notes_text_frame.text = (
        "First term drains the backlog inside the latency target; second keeps up with steady arrivals. Ten HED "
        "tiles and ten Canny tiles look identical to KEDA, very different to us.")

    # 12 ── autoscaler results
    s = slide_base(prs, "HPA vs KEDA vs work-aware on the same replayed trace", "Upgrade 4 · Results", "4")
    if {"hpa", "keda", "work"} <= auto.keys():
        names = ["HPA", "KEDA", "Work-aware"]
        keys = ["hpa", "keda", "work"]
        for i, (metric, title, fmt) in enumerate((("p95_s", "p95 latency (s)", '0.0'),
                                                  ("within_target_pct", "jobs within 30 s target (%)", '0'),
                                                  ("pod_seconds", "worker pod-seconds (cost)", '0'))):
            cd = CategoryChartData()
            cd.categories = names
            cd.add_series(title, [auto[k][metric] for k in keys])
            gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5 + 4.15 * i), Inches(1.55),
                                    Inches(4.0), Inches(3.6), cd)
            ch = gf.chart
            style_chart(ch, title, legend=False, number_format=fmt)
            ch.plots[0].gap_width = 70
            ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            ch.plots[0].vary_by_categories = False
            ser = ch.plots[0].series[0]
            for j, c in enumerate([GREY, AMBER, TEAL]):
                pt = ser.points[j]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = c
        a = next(iter(auto.values()))
        text(s, Inches(0.6), Inches(5.35), Inches(12.2), Inches(1.6),
             autoscale_takeaway(auto) + f"  (Minikube, 8 CPUs; {a['jobs']} jobs: cheap and expensive bursts; "
             f"mean of {a['runs']} runs each.)", size=15)
    else:
        text(s, Inches(0.6), Inches(2), Inches(12), Inches(1), "Run scripts/run_autoscale_comparison.sh first.")
    s.notes_slide.notes_text_frame.text = (
        "Same seeded trace for all three, interleaved runs. Discuss latency vs cost; HPA reacts late because "
        "CPU is averaged over a window, KEDA reacts to tile counts regardless of cost.")

    # 13 ── deployment + testing
    s = slide_base(prs, "Laptop, Compose, Kubernetes — and the path to AWS", "Deployment & quality")
    rows = [("Object store", "MinIO", "Amazon S3 + KMS"), ("Queue / state", "Redis", "ElastiCache"),
            ("Orchestration", "Minikube", "Amazon EKS"), ("Pod scaling", "Work-aware / HPA / KEDA",
                                                         "same + Karpenter"),
            ("Credentials", "K8s Secret", "IRSA roles"), ("Images", "Minikube Docker", "Amazon ECR")]
    tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(0.6), Inches(1.6), Inches(7.4), Inches(4.2)).table
    for j, h in enumerate(("Component", "Ran here", "On AWS")):
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        r = c.text_frame.paragraphs[0].runs[0]
        r.font.color.rgb = WHITE
        r.font.bold = True
        r.font.size = Pt(15)
        r.font.name = BODY
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = tbl.cell(i, j)
            c.text = v
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xF1, 0xF5, 0xF9)
            r = c.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(14)
            r.font.name = BODY
            r.font.color.rgb = INK
    stat(s, Inches(8.6), Inches(1.6), Inches(4.3), "43 tests", "all pass · 17 new for the upgrades", vsize=36)
    bullets(s, Inches(8.6), Inches(3.2), Inches(4.3), Inches(3.6), [
        "Seam-free:|tiled = whole image (classical exact; neural ≥ 40 dB seams)",
        "Fault tolerance kept:|killed worker → leased tiles re-run",
        "Exact accounting:|work counter survives retries and duplicates",
        "AWS templates:|written, not deployed",
    ], size=14, gap=6)
    s.notes_slide.notes_text_frame.text = (
        "k8s manifests include MinIO with NodePort for presigned downloads, Secret for credentials, RBAC for "
        "the autoscaler. deploy/aws has bucket policy, IAM, lifecycle and IRSA - not executed.")

    # 14 ── limitations
    s = slide_base(prs, "What we learned, and what is left", "Discussion")
    card(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.1), "Lessons",
         "Whether and how finely to distribute matters as much as the parallel machinery.\n\n"
         "Object storage trades ~1.3 ms per tile for removing the memory ceiling.\n\n"
         "Effective receptive fields are far smaller than theoretical ones.\n\n"
         "Neural ops are memory-bound on one host: measure, don't assume.", size=15)
    card(s, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.1), "Limitations & future work",
         "Single host: a multi-node EKS run would test the model on real networks.\n\n"
         "Cheap ops should run in-process at the API, not through the queue.\n\n"
         "Not built: GPU workers, spot node pools, COG range reads + deep-zoom viewer, learned cost models.\n\n"
         "MAPE 20–25 %: good for decisions, not SLAs.", accent=AMBER, size=15)
    s.notes_slide.notes_text_frame.text = "Be upfront about the single-host setup and the unexecuted AWS templates."

    # 15 ── close
    s = slide_base(prs, "", dark=True)
    text(s, Inches(0.8), Inches(1.2), Inches(11.5), Inches(1.0), "From tile-parallel to cost-aware",
         size=40, font=HEAD, color=WHITE, bold=True)
    pts = [("①", f"{rm['redis_bytes_per_8192_job'] / 1e6:.0f} MB → {sm['redis_bytes_per_8192_job'] / 1e3:.0f} KB of Redis per job"),
           ("②", f"HED + Real-ESRGAN, halos {halo['hed']['halo']} px / {halo['superres']['halo']} px, seam-free"),
           ("③", f"{eng['mape_pct']:.0f} % prediction error; no more slowdowns from naive tiling"),
           ("④", "autoscaling on queued compute-seconds, compared with HPA and KEDA")]
    for i, (n, t) in enumerate(pts):
        y = Inches(2.6 + 0.8 * i)
        text(s, Inches(0.8), y, Inches(0.6), Inches(0.6), n, size=26, color=TEAL, bold=True)
        text(s, Inches(1.5), y + Inches(0.05), Inches(11), Inches(0.6), t, size=21, color=WHITE)
    text(s, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.5),
         "Thank you  ·  github.com/GHOSTALERT/cloud-native-ing-proc", size=16, color=RGBColor(0xCB, 0xD5, 0xE1))
    s.notes_slide.notes_text_frame.text = "Summarise the four upgrades in one line each and invite questions."

    out = ROOT / "docs" / "Project_Presentation.pptx"
    prs.save(out)
    return out


def autoscale_takeaway(auto):
    h, k, w = auto["hpa"], auto["keda"], auto["work"]
    return (f"HPA is cheapest but reacts after {h['first_scale_up_s']:.0f} s and misses the target. KEDA and "
            f"work-aware meet it for every job; work-aware uses {100 * (1 - w['pod_seconds'] / k['pod_seconds']):.0f} % "
            "fewer pod-seconds because it ignores bursts of cheap tiles (KEDA scales up within 4 s of them), with "
            f"p95 {w['p95_s']:.1f} s vs {k['p95_s']:.1f} s. Dropping to 1 worker 10 s after the queue empties saves a further "
            "~12 % over our first version.")

if __name__ == "__main__":
    print("wrote", build())
