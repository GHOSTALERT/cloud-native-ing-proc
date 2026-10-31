"""Small python-docx toolkit shared by the report (IEEE layout) and the run guide.

Markup in text: **bold**, *italic*, `code`, ^{superscript}, _{subscript}.
"""
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV"]
MARKUP = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|\^\{.+?\}|_\{.+?\})")


def _shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


class Doc:
    def __init__(self, font="Times New Roman", size=10, margins_cm=(1.9, 1.6, 2.5, 1.6),
                 code_font="Courier New", accent=RGBColor(0, 0, 0)):
        self.d = Document()
        self.font, self.size, self.code_font, self.accent = font, size, code_font, accent
        self.fig = self.tab = self.eq = 0
        self.h1 = self.h2 = 0
        sec = self.d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin, sec.left_margin, sec.bottom_margin, sec.right_margin = (Cm(m) for m in margins_cm)
        st = self.d.styles["Normal"]
        st.font.name = font
        st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        st.font.size = Pt(size)
        pf = st.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
            h = self.d.styles[name]
            h.font.name = font
            rpr = h.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.append(fonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(attr), font)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if fonts.get(qn(attr)) is not None:
                    del fonts.attrib[qn(attr)]
            h.font.color.rgb = accent
            h.font.italic = False

    # ------------------------------------------------------------ layout
    @property
    def section(self):
        return self.d.sections[-1]

    def width_in(self, span=False):
        sec = self.section
        full = (sec.page_width - sec.left_margin - sec.right_margin) / 914400
        cols = sec._sectPr.find(qn("w:cols"))
        n = int(cols.get(qn("w:num"), "1")) if cols is not None else 1
        if n == 1 or span:
            return full
        space = int(cols.get(qn("w:space"), "0")) / 1440
        return (full - space * (n - 1)) / n

    def columns(self, n, space_cm=0.42, new_page=False):
        """Start a new section with n text columns (continuous unless new_page)."""
        sec = self.d.add_section(WD_SECTION.NEW_PAGE if new_page else WD_SECTION.CONTINUOUS)
        cols = sec._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sec._sectPr.append(cols)
        cols.set(qn("w:num"), str(n))
        cols.set(qn("w:space"), str(int(space_cm / 2.54 * 1440)))
        return sec

    def page_numbers(self, size=8):
        for sec in self.d.sections:
            sec.footer.is_linked_to_previous = True
        p = self.d.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        run.font.size = Pt(size)

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def column_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)

    # ------------------------------------------------------------- text
    def rich(self, para, text, bold=False, italic=False, size=None, color=None, font=None, caps=False):
        for part in MARKUP.split(text):
            if not part:
                continue
            sup = sub = False
            if part.startswith("**") and part.endswith("**"):
                r = para.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = para.add_run(part[1:-1])
                r.italic = True
            elif part.startswith("`"):
                r = para.add_run(part[1:-1])
                r.font.name = self.code_font
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), self.code_font)
                r.font.size = Pt((size or self.size) - 1)
            elif part.startswith("^{"):
                r = para.add_run(part[2:-1])
                sup = True
            elif part.startswith("_{"):
                r = para.add_run(part[2:-1])
                sub = True
            else:
                r = para.add_run(part)
            if bold:
                r.bold = True
            if italic:
                r.italic = True
            if size and not part.startswith("`"):
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
            if font:
                r.font.name = font
            if caps:
                r.font.small_caps = True
            r.font.superscript = sup or None
            r.font.subscript = sub or None
        return para

    def p(self, text="", *, bold=False, italic=False, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          indent=True, before=0, after=0, color=None, keep=False):
        para = self.d.add_paragraph()
        self.rich(para, text, bold, italic, size, color)
        pf = para.paragraph_format
        pf.alignment = align
        if indent is True:
            pf.first_line_indent = Cm(0.36)
        elif indent:
            pf.first_line_indent = Cm(indent)
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.keep_with_next = keep or None
        return para

    def bullets(self, items, size=None, bullet="•", indent_cm=0.5, after=1):
        for it in items:
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(indent_cm)
            pf.first_line_indent = Cm(-0.3)
            pf.space_after = Pt(after)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.tab_stops.add_tab_stop(Cm(indent_cm))
            self.rich(para, f"{bullet}\t{it}", size=size)

    def numbered(self, items, size=None, after=1, start=1):
        for i, it in enumerate(items, start):
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(0.6)
            pf.first_line_indent = Cm(-0.45)
            pf.space_after = Pt(after)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.tab_stops.add_tab_stop(Cm(0.6))
            self.rich(para, f"{i})\t{it}", size=size)

    def code(self, text, size=7.5, fill="F3F4F6"):
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.space_before = Pt(3) if i == 0 else Pt(0)
            pf.space_after = Pt(3) if i == len(lines) - 1 else Pt(0)
            pf.left_indent = Cm(0.15)
            pf.keep_with_next = i < len(lines) - 1 or None
            r = para.add_run(line or " ")
            r.font.name = self.code_font
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), self.code_font)
            r.font.size = Pt(size)
            _shade(para._p.get_or_add_pPr(), fill)

    def box(self, text, fill="EFF6FF", label=None, size=None):
        para = self.d.add_paragraph()
        self.rich(para, (f"**{label}** " if label else "") + text, size=size)
        pf = para.paragraph_format
        pf.left_indent = pf.right_indent = Cm(0.15)
        pf.space_before = pf.space_after = Pt(4)
        _shade(para._p.get_or_add_pPr(), fill)
        return para

    def equation(self, text, size=None):
        """Display equation, centred, numbered (n) at the right margin."""
        self.eq += 1
        para = self.d.add_paragraph()
        pf = para.paragraph_format
        w = self.width_in()
        pf.tab_stops.add_tab_stop(Inches(w / 2), alignment=1)   # centre
        pf.tab_stops.add_tab_stop(Inches(w), alignment=2)       # right
        pf.space_before = pf.space_after = Pt(3)
        self.rich(para, "\t" + text, italic=False, size=size)
        para.add_run(f"\t({self.eq})")
        return self.eq

    # --------------------------------------------------- IEEE headings
    def ieee_h1(self, text):
        self.h1 += 1
        self.h2 = 0
        para = self.d.add_paragraph()
        pf = para.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before, pf.space_after = Pt(8), Pt(4)
        pf.keep_with_next = True
        r = para.add_run(f"{ROMAN[self.h1 - 1]}. {text}")
        r.font.small_caps = True
        r.font.size = Pt(self.size)
        return para

    def ieee_h2(self, text):
        self.h2 += 1
        para = self.d.add_paragraph()
        pf = para.paragraph_format
        pf.space_before, pf.space_after = Pt(5), Pt(2)
        pf.keep_with_next = True
        r = para.add_run(f"{chr(64 + self.h2)}. {text}")
        r.italic = True
        r.font.size = Pt(self.size)
        return para

    def h(self, text, level=1, size=None, before=10, after=4, color=None):
        para = self.d.add_paragraph(style=f"Heading {level}")
        para.paragraph_format.space_before, para.paragraph_format.space_after = Pt(before), Pt(after)
        r = para.add_run(text)
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        return para

    # ------------------------------------------------- tables / figures
    def table(self, header, rows, caption=None, widths_cm=None, size=8, ieee=True, header_fill=None,
              span=False, align_cols=None, zebra=None):
        """IEEE style: caption above ('TABLE I' + small caps title), rules only."""
        if caption:
            self.tab += 1
            if ieee:
                c1 = self.d.add_paragraph()
                c1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                c1.paragraph_format.space_before = Pt(6)
                c1.paragraph_format.keep_with_next = True
                c1.add_run(f"TABLE {ROMAN[self.tab - 1]}").font.size = Pt(size)
                c2 = self.d.add_paragraph()
                c2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                c2.paragraph_format.space_after = Pt(3)
                c2.paragraph_format.keep_with_next = True
                r = c2.add_run(caption)
                r.font.small_caps = True
                r.font.size = Pt(size)
            else:
                c = self.p(f"**Table {self.tab}.** {caption}", size=size + 0.5, indent=False,
                           align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3, keep=True)
        t = self.d.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if not ieee:
            t.style = "Table Grid"
        widths = widths_cm or [self.width_in(span) * 2.54 / len(header)] * len(header)
        tblpr = t._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblpr.append(layout)
        if ieee:
            borders = OxmlElement("w:tblBorders")
            for edge, sz in (("top", 8), ("bottom", 8), ("insideH", 0), ("insideV", 0), ("left", 0), ("right", 0)):
                b = OxmlElement(f"w:{edge}")
                b.set(qn("w:val"), "single" if sz else "nil")
                if sz:
                    b.set(qn("w:sz"), str(sz))
                    b.set(qn("w:color"), "000000")
                borders.append(b)
            tblpr.append(borders)

        def fill_row(cells, values, bold=False, is_header=False):
            for i, v in enumerate(values):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.alignment = (align_cols[i] if align_cols else
                                                   (WD_ALIGN_PARAGRAPH.CENTER if is_header or i else WD_ALIGN_PARAGRAPH.LEFT))
                color = RGBColor(255, 255, 255) if (is_header and header_fill) else None
                self.rich(para, str(v), bold=bold, size=size, color=color)
                if is_header and header_fill:
                    _shade(cells[i]._tc.get_or_add_tcPr(), header_fill)

        fill_row(t.rows[0].cells, header, bold=True, is_header=True)
        if ieee:  # rule under the header
            for c in t.rows[0].cells:
                tcpr = c._tc.get_or_add_tcPr()
                b = OxmlElement("w:tcBorders")
                bot = OxmlElement("w:bottom")
                bot.set(qn("w:val"), "single")
                bot.set(qn("w:sz"), "4")
                bot.set(qn("w:color"), "000000")
                b.append(bot)
                tcpr.append(b)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            bold = isinstance(row, tuple) and len(row) and str(row[0]).startswith("**")
            fill_row(cells, [str(x).replace("**", "") if bold else x for x in row], bold=bold)
            if zebra and ri % 2:
                for c in cells:
                    _shade(c._tc.get_or_add_tcPr(), zebra)
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
            for row in t.rows:
                row.cells[i].width = Cm(w)
        gap = self.d.add_paragraph()
        gap.paragraph_format.space_after = Pt(4)
        return t

    def figure(self, path: Path, caption: str, width_in=None, max_h_in=4.2, ieee=True, size=8):
        with Image.open(path) as im:
            w_px, h_px = im.size
        width = width_in or self.width_in()
        if width * h_px / w_px > max_h_in:
            width = max_h_in * w_px / h_px
        para = self.d.add_paragraph()
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Inches(width))
        self.fig += 1
        cap = self.d.add_paragraph()
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if ieee else WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(6)
        self.rich(cap, (f"Fig. {self.fig}. " if ieee else f"**Figure {self.fig}.** ") + caption, size=size)
        return self.fig

    def references(self, refs, size=8):
        for i, ref in enumerate(refs, 1):
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(0.65)
            pf.first_line_indent = Cm(-0.65)
            pf.space_after = Pt(1.5)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.tab_stops.add_tab_stop(Cm(0.65))
            self.rich(para, f"[{i}]\t{ref}", size=size)

    def save(self, path: Path):
        self.d.save(path)
        return path


def to_pdf(paths, outdir: Path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found; skipping PDF export")
        return
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), *map(str, paths)],
                   check=True, capture_output=True, timeout=600)
